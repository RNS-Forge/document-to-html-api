"""
Document-to-HTML API — High-Accuracy Universal Document Converter
==================================================================
Converts PDF, DOCX, PPTX, XLSX, images (PNG/JPG/WEBP/BMP/TIFF/GIF),
SVG, Markdown, CSV, and plain text into high-fidelity, semantic HTML.

Key Features:
  • ALL images embedded as Base64 data URIs (100% self-contained HTML).
  • PDF: Visual reading order sorting (y0, x0), table detection with bounding
    box exclusion, font-aware heading detection, color preservation,
    hyperlinks, and high-res OCR fallback for scanned pages.
  • DOCX: Mammoth semantic mapping + python-docx fallback, embedded Base64 images.
  • PPTX: Slide cards, recursive group shape processing, picture shapes to Base64,
    per-run typography and hyperlinks, table extraction, and speaker notes.
  • XLSX: Multi-sheet tab navigation, merged cells, font & background fill colors,
    date/number formatting, and embedded sheet images to Base64.
  • Images: High-res Base64 rendering, metadata pills, and intelligent OCR.
  • Local Test UI: Drag-and-drop web preview with live iframe, source code viewer,
    download option, and conversion metrics (when LOCAL_RUN=true).
"""

import io
import os
import csv
import json
import base64
import time
import shutil
import traceback
from pathlib import Path
from typing import Optional, List, Dict, Any, Tuple

from dotenv import load_dotenv
load_dotenv()

from fastapi import FastAPI, File, UploadFile, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, HTMLResponse

# ─── Configuration ────────────────────────────────────────────────────────────

LOCAL_RUN = os.getenv("LOCAL_RUN", "true").strip().lower() == "true"
PORT = int(os.getenv("PORT", "8000"))
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", 50 * 1024 * 1024))  # 50 MB default
TESSERACT_CMD_ENV = os.getenv("TESSERACT_CMD", "").strip()

# ─── Tesseract Setup & Detection ──────────────────────────────────────────────

TESSERACT_AVAILABLE = False
try:
    import pytesseract
    # Auto-detect common Windows paths if not configured
    candidates = [
        TESSERACT_CMD_ENV,
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
    ]
    for cand in candidates:
        if cand and os.path.exists(cand):
            pytesseract.pytesseract.tesseract_cmd = cand
            break

    # Verify if tesseract actually runs
    try:
        pytesseract.get_tesseract_version()
        TESSERACT_AVAILABLE = True
    except Exception:
        TESSERACT_AVAILABLE = False
except ImportError:
    TESSERACT_AVAILABLE = False

# ─── App Setup ────────────────────────────────────────────────────────────────

app = FastAPI(
    title="Document-to-HTML API",
    description="Universal document converter producing pixel-accurate, self-contained semantic HTML.",
    version="2.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ─── Extension → Handler Mapping ─────────────────────────────────────────────

EXTENSION_MAP = {
    ".pdf": "pdf",
    ".docx": "docx", ".doc": "docx",
    ".pptx": "pptx", ".ppt": "pptx",
    ".xlsx": "xlsx", ".xls": "xlsx",
    ".png": "image", ".jpg": "image", ".jpeg": "image",
    ".webp": "image", ".bmp": "image", ".tiff": "image",
    ".tif": "image", ".gif": "image", ".ico": "image",
    ".svg": "svg",
    ".md": "markdown", ".markdown": "markdown",
    ".csv": "csv", ".tsv": "csv",
    ".json": "json",
    ".xml": "xml",
    ".txt": "text", ".log": "text", ".yaml": "text", ".yml": "text",
    ".html": "passthrough", ".htm": "passthrough",
}

# ─── Shared Responsive CSS ───────────────────────────────────────────────────

SHARED_CSS = """
<style>
  :root {
    --doc-font: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
    --doc-mono: 'Cascadia Code', 'Fira Code', 'Consolas', monospace;
    --doc-text: #1e293b;
    --doc-bg: #ffffff;
    --doc-border: #e2e8f0;
    --doc-primary: #3b82f6;
  }
  *, *::before, *::after { box-sizing: border-box; }
  .doc-root {
    font-family: var(--doc-font);
    color: var(--doc-text);
    background: var(--doc-bg);
    line-height: 1.65;
    padding: 32px 40px;
    max-width: 960px;
    margin: 0 auto;
    word-wrap: break-word;
    font-size: 15px;
  }
  .doc-root h1, .doc-root h2, .doc-root h3, .doc-root h4, .doc-root h5, .doc-root h6 {
    color: #0f172a;
    font-weight: 700;
    line-height: 1.3;
    margin: 1.1em 0 0.4em;
  }
  .doc-root h1 { font-size: 1.9em; border-bottom: 2px solid var(--doc-border); padding-bottom: 6px; }
  .doc-root h2 { font-size: 1.45em; border-bottom: 1px solid var(--doc-border); padding-bottom: 4px; }
  .doc-root h3 { font-size: 1.2em; }
  .doc-root h4 { font-size: 1.05em; }
  .doc-root p { margin: 0.5em 0; }
  .doc-root a { color: var(--doc-primary); text-decoration: underline; text-underline-offset: 2px; }
  .doc-root a:hover { color: #1d4ed8; }
  .doc-root ul, .doc-root ol { margin: 0.5em 0; padding-left: 1.8em; }
  .doc-root li { margin-bottom: 0.25em; }
  .doc-root table {
    border-collapse: collapse;
    width: 100%;
    margin: 14px 0;
    font-size: 0.93em;
    overflow-x: auto;
    display: table;
  }
  .doc-table-wrap { overflow-x: auto; margin: 14px 0; border: 1px solid var(--doc-border); border-radius: 6px; }
  .doc-table-wrap table { margin: 0; border: none; }
  .doc-root th, .doc-root td {
    border: 1px solid var(--doc-border);
    padding: 8px 12px;
    text-align: left;
    vertical-align: top;
  }
  .doc-root th {
    background: #f8fafc;
    font-weight: 600;
    color: #0f172a;
  }
  .doc-root tr:nth-child(even) td { background: #fafafa; }
  .doc-root tr:hover td { background: #f1f5f9; }
  .doc-root img {
    max-width: 100%;
    height: auto;
    display: inline-block;
    border-radius: 4px;
    margin: 8px 0;
  }
  .doc-img-wrap {
    text-align: center;
    margin: 14px 0;
  }
  .doc-img-wrap img {
    box-shadow: 0 1px 4px rgba(0,0,0,0.08);
  }
  .doc-root pre {
    background: #0f172a;
    color: #f8fafc;
    padding: 14px 18px;
    border-radius: 8px;
    overflow-x: auto;
    font-size: 0.88em;
    line-height: 1.5;
    margin: 12px 0;
  }
  .doc-root code {
    font-family: var(--doc-mono);
    font-size: 0.9em;
  }
  .doc-root p code, .doc-root li code {
    background: #f1f5f9;
    color: #0f172a;
    padding: 2px 5px;
    border-radius: 4px;
    border: 1px solid var(--doc-border);
  }
  .doc-root blockquote {
    border-left: 4px solid var(--doc-primary);
    padding: 8px 16px;
    margin: 12px 0;
    background: #eff6ff;
    color: #1e3a8a;
    border-radius: 0 4px 4px 0;
  }
  .doc-root hr { border: none; border-top: 1.5px solid var(--doc-border); margin: 20px 0; }
  .page-box {
    margin-bottom: 24px;
    padding-bottom: 24px;
    border-bottom: 2px dashed #cbd5e1;
  }
  .page-box:last-child { border-bottom: none; }
  .page-label {
    font-size: 0.75rem;
    font-weight: 600;
    color: #64748b;
    text-transform: uppercase;
    letter-spacing: 0.8px;
    margin-bottom: 10px;
    display: flex;
    align-items: center;
    gap: 8px;
  }
  .page-label::after { content: ''; flex: 1; height: 1px; background: #e2e8f0; }
  .slide-card {
    border: 1px solid var(--doc-border);
    border-radius: 12px;
    padding: 24px 28px;
    margin-bottom: 22px;
    background: #ffffff;
    box-shadow: 0 2px 6px rgba(0,0,0,0.04);
  }
  .slide-title {
    font-size: 1.35em;
    font-weight: 700;
    color: #0f172a;
    margin-bottom: 12px;
    border-bottom: 2px solid #6366f1;
    padding-bottom: 6px;
  }
  .notes-drawer {
    margin-top: 16px;
    background: #f8fafc;
    border: 1px solid var(--doc-border);
    border-radius: 6px;
    padding: 8px 12px;
    font-size: 0.88em;
  }
  .notes-drawer summary { cursor: pointer; font-weight: 600; color: #475569; }
  .notes-body { margin-top: 8px; color: #334155; }
  .meta-badge-bar {
    display: flex;
    flex-wrap: wrap;
    gap: 8px;
    margin-bottom: 16px;
    font-size: 0.78rem;
  }
  .meta-badge {
    background: #f1f5f9;
    border: 1px solid #cbd5e1;
    padding: 3px 8px;
    border-radius: 4px;
    color: #334155;
    font-weight: 500;
  }
  .ocr-box {
    margin-top: 14px;
    padding: 14px 18px;
    background: #fffbeb;
    border: 1px solid #fde68a;
    border-radius: 8px;
    font-size: 0.9em;
    color: #78350f;
  }
  .ocr-title {
    font-weight: 700;
    font-size: 0.85em;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    margin-bottom: 6px;
    color: #b45309;
  }
  .tab-container { margin: 16px 0; }
  .tab-buttons { display: flex; flex-wrap: wrap; gap: 4px; border-bottom: 2px solid #e2e8f0; padding-bottom: 0; }
  .tab-btn-sheet {
    padding: 8px 16px;
    border: 1px solid #cbd5e1;
    border-bottom: none;
    border-radius: 6px 6px 0 0;
    background: #f1f5f9;
    color: #475569;
    font-weight: 600;
    font-size: 0.88em;
    cursor: pointer;
  }
  .tab-btn-sheet.active { background: #ffffff; color: var(--doc-primary); border-top: 2px solid var(--doc-primary); }
  .tab-panel-sheet { display: none; padding: 16px 0; }
  .tab-panel-sheet.active { display: block; }

  @media print {
    .doc-root { padding: 0; max-width: 100%; }
    .page-box { page-break-after: always; border: none; margin: 0; }
    .slide-card { page-break-after: always; box-shadow: none; border: none; }
  }
</style>
""".strip()


# ─── General Helpers ──────────────────────────────────────────────────────────

def _esc(text: Any) -> str:
    """HTML-escape a string cleanly."""
    if text is None:
        return ""
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )


def _to_base64_data_uri(img_bytes: bytes, mime_type: str = "image/png") -> str:
    """Return a complete base64 data URI string."""
    b64 = base64.b64encode(img_bytes).decode("utf-8")
    return f"data:{mime_type};base64,{b64}"


def _detect_img_mime(data: bytes, filename: str = "") -> str:
    """Detect image MIME type from magic numbers or filename."""
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data.startswith(b"GIF87a") or data.startswith(b"GIF89a"):
        return "image/gif"
    if data.startswith(b"RIFF") and len(data) > 12 and data[8:12] == b"WEBP":
        return "image/webp"
    if data.startswith(b"BM"):
        return "image/bmp"
    if data.startswith(b"II*\x00") or data.startswith(b"MM\x00*"):
        return "image/tiff"
    if data.startswith(b"<svg") or b"<svg" in data[:200]:
        return "image/svg+xml"

    ext = Path(filename).suffix.lower().lstrip(".")
    mime_map = {
        "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
        "webp": "image/webp", "bmp": "image/bmp", "gif": "image/gif",
        "tiff": "image/tiff", "tif": "image/tiff", "svg": "image/svg+xml",
        "ico": "image/x-icon",
    }
    return mime_map.get(ext, "image/png")


def wrap_html(body: str, title: str = "Converted Document") -> str:
    """Wrap inner HTML in a complete, self-contained HTML5 document."""
    return (
        f"<!DOCTYPE html>\n<html lang=\"en\">\n<head>\n"
        f"  <meta charset=\"UTF-8\">\n"
        f"  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\">\n"
        f"  <title>{_esc(title)}</title>\n"
        f"  {SHARED_CSS}\n"
        f"</head>\n<body>\n"
        f"  <div class=\"doc-root\">\n{body}\n  </div>\n"
        f"</body>\n</html>"
    )


def _run_ocr_image(img_bytes: bytes) -> str:
    """Safely run OCR on image bytes and return clean extracted text."""
    if not TESSERACT_AVAILABLE:
        return ""
    try:
        from PIL import Image, ImageEnhance
        img = Image.open(io.BytesIO(img_bytes))
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")

        # Contrast enhancement for sharper OCR recognition
        if img.width < 1200:
            scale = min(2.0, 1600 / max(img.width, 1))
            new_size = (int(img.width * scale), int(img.height * scale))
            img = img.resize(new_size, Image.Resampling.LANCZOS)

        enhancer = ImageEnhance.Contrast(img)
        img = enhancer.enhance(1.4)

        text = pytesseract.image_to_string(img, lang="eng", config="--psm 3")
        return text.strip()
    except Exception:
        return ""


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# CONVERTERS
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━


# ── 1. PDF Converter (High-Accuracy Engine) ──────────────────────────────────

def convert_pdf(data: bytes, filename: str) -> str:
    import pymupdf

    doc = pymupdf.open(stream=data, filetype="pdf")
    total_pages = len(doc)
    pages_html = []

    # Calculate global font size distribution to infer body text size
    font_sizes = []
    for p in doc:
        td = p.get_text("dict", flags=0)
        for b in td.get("blocks", []):
            if b.get("type") == 0:
                for ln in b.get("lines", []):
                    for sp in ln.get("spans", []):
                        sz = sp.get("size", 0)
                        if sz > 0:
                            font_sizes.append(sz)

    median_body_size = 11.0
    if font_sizes:
        font_sizes.sort()
        median_body_size = font_sizes[len(font_sizes) // 2]
    if median_body_size < 8:
        median_body_size = 11.0

    for pg_num in range(total_pages):
        page = doc[pg_num]
        page_rect = page.rect
        page_width = page_rect.width
        elements = []  # List of dicts: {"y0": float, "x0": float, "type": str, "html": str}

        # ── Step 1: Extract Hyperlinks ──
        page_links = []
        try:
            links = page.get_links()
            for lk in links:
                uri = lk.get("uri")
                lk_rect = lk.get("from")
                if uri and lk_rect:
                    page_links.append({"rect": pymupdf.Rect(lk_rect), "uri": uri})
        except Exception:
            pass

        # ── Step 2: Extract Tables ──
        table_rects = []
        try:
            tab_finder = page.find_tables()
            for tab in tab_finder.tables:
                tab_rect = pymupdf.Rect(tab.bbox)
                table_rects.append(tab_rect)

                # Render structured HTML table
                extracted = tab.extract()
                if extracted and len(extracted) > 0:
                    tbl_html = _render_table_data(extracted)
                    wrapped_tbl = f'<div class="doc-table-wrap">{tbl_html}</div>'
                    elements.append({
                        "y0": tab_rect.y0,
                        "x0": tab_rect.x0,
                        "type": "table",
                        "html": wrapped_tbl,
                    })
        except Exception:
            pass

        # ── Step 3: Extract Embedded Images as Base64 ──
        extracted_xrefs = set()
        image_rects = []
        try:
            img_infos = page.get_image_info(xrefs=True)
            for info in img_infos:
                xref = info.get("xref", 0)
                bbox_t = info.get("bbox", (0, 0, 0, 0))
                img_rect = pymupdf.Rect(bbox_t)

                # Skip zero-sized or tiny 1px separator lines
                if img_rect.width < 8 or img_rect.height < 8:
                    continue

                image_rects.append(img_rect)
                if xref > 0 and xref not in extracted_xrefs:
                    extracted_xrefs.add(xref)
                    base_img = doc.extract_image(xref)
                    if base_img:
                        img_bytes = base_img["image"]
                        ext = base_img.get("ext", "png")
                        mime = f"image/{ext}" if ext != "jpg" else "image/jpeg"
                        data_uri = _to_base64_data_uri(img_bytes, mime)
                        img_html = (
                            f'<div class="doc-img-wrap">'
                            f'  <img src="{data_uri}" alt="Embedded Figure" '
                            f'       style="max-width:100%;height:auto;">'
                            f'</div>'
                        )
                        elements.append({
                            "y0": img_rect.y0,
                            "x0": img_rect.x0,
                            "type": "image",
                            "html": img_html,
                        })
        except Exception:
            pass

        # ── Step 4: Extract Text Blocks (dict) ──
        pdf_flags = pymupdf.TEXT_PRESERVE_WHITESPACE | pymupdf.TEXT_PRESERVE_SPANS | pymupdf.TEXT_DEHYPHENATE
        text_dict = page.get_text("dict", flags=pdf_flags)
        blocks = text_dict.get("blocks", [])

        has_text = False
        for block in blocks:
            if block.get("type") == 0:  # Text block
                block_rect = pymupdf.Rect(block["bbox"])

                # Skip text that significantly overlaps with detected tables
                if any(block_rect.intersects(tr) for tr in table_rects):
                    overlap = False
                    for tr in table_rects:
                        inter = block_rect & tr
                        if inter.width * inter.height > 0.4 * (block_rect.width * block_rect.height):
                            overlap = True
                            break
                    if overlap:
                        continue

                block_html = _render_pdf_text_block(block, median_body_size, page_width, page_links)
                if block_html.strip():
                    has_text = True
                    elements.append({
                        "y0": block_rect.y0,
                        "x0": block_rect.x0,
                        "type": "text",
                        "html": block_html,
                    })

        # ── Step 5: Visual Reading Order Sort ──
        elements.sort(key=lambda el: (round(el["y0"] / 6.0) * 6.0, el["x0"]))

        # Consecutive list normalization (group adjacent <li> into <ul>)
        page_parts = []
        list_buffer = []

        def flush_list():
            if list_buffer:
                page_parts.append(f"<ul>\n  " + "\n  ".join(list_buffer) + "\n</ul>")
                list_buffer.clear()

        for el in elements:
            html = el["html"]
            if html.startswith("<li>") and html.endswith("</li>"):
                list_buffer.append(html)
            else:
                flush_list()
                page_parts.append(html)
        flush_list()

        # ── Step 6: Scanned Page Fallback (Raster Pixmap + OCR) ──
        if not has_text and not table_rects:
            try:
                mat = pymupdf.Matrix(2.0, 2.0)
                pix = page.get_pixmap(matrix=mat)
                img_bytes = pix.tobytes("png")
                data_uri = _to_base64_data_uri(img_bytes, "image/png")

                fallback_parts = [
                    f'<div class="doc-img-wrap">'
                    f'  <img src="{data_uri}" alt="Scanned Page {pg_num+1}">'
                    f'</div>'
                ]

                ocr_text = _run_ocr_image(img_bytes)
                if ocr_text:
                    ocr_lines = [_esc(l) for l in ocr_text.split("\n") if l.strip()]
                    fallback_parts.append(
                        f'<div class="ocr-box">'
                        f'  <div class="ocr-title">Recognized Text (OCR)</div>'
                        f'  <div>{"<br>".join(ocr_lines)}</div>'
                        f'</div>'
                    )

                page_parts = fallback_parts
            except Exception:
                pass

        # Assemble page container
        page_content = "\n    ".join(page_parts) if page_parts else "<p><em>(Blank page)</em></p>"
        if total_pages > 1:
            pages_html.append(
                f'<section class="page-box" data-page="{pg_num+1}">\n'
                f'  <div class="page-label">Page {pg_num+1} of {total_pages}</div>\n'
                f'  {page_content}\n'
                f'</section>'
            )
        else:
            pages_html.append(page_content)

    doc.close()
    inner = "\n".join(pages_html) if pages_html else "<p><em>No extractable content found in PDF.</em></p>"
    return wrap_html(inner, filename)


def _render_pdf_text_block(block: dict, median_body_size: float, page_width: float, page_links: list) -> str:
    """Render a text block preserving typography, colors, links, and alignments."""
    import pymupdf
    lines_html = []
    max_font_size = 0.0

    block_x0 = block.get("bbox", [0, 0, 0, 0])[0]
    block_x1 = block.get("bbox", [0, 0, 0, 0])[2]
    block_width = block_x1 - block_x0

    align_style = ""
    center_offset = abs((page_width - block_width) / 2.0 - block_x0)
    if center_offset < 24 and block_width < page_width * 0.75:
        align_style = ' style="text-align: center;"'
    elif block_x0 > page_width * 0.55 and (page_width - block_x1) < 60:
        align_style = ' style="text-align: right;"'

    for line in block.get("lines", []):
        spans_html = []
        for span in line.get("spans", []):
            text = span.get("text", "")
            if not text:
                continue

            size = span.get("size", 11.0)
            flags = span.get("flags", 0)
            color_int = span.get("color", 0)
            font_name = span.get("font", "")
            span_bbox = pymupdf.Rect(span.get("bbox", (0, 0, 0, 0)))

            max_font_size = max(max_font_size, size)

            is_bold = bool(flags & 16) or any(k in font_name.lower() for k in ("bold", "black", "heavy", "medium"))
            is_italic = bool(flags & 2) or any(k in font_name.lower() for k in ("italic", "oblique"))
            is_mono = any(k in font_name.lower() for k in ("mono", "courier", "code", "consolas"))
            is_superscript = bool(flags & 1)

            escaped = _esc(text)

            if is_mono:
                escaped = f"<code>{escaped}</code>"
            if is_bold and is_italic:
                escaped = f"<strong><em>{escaped}</em></strong>"
            elif is_bold:
                escaped = f"<strong>{escaped}</strong>"
            elif is_italic:
                escaped = f"<em>{escaped}</em>"
            if is_superscript:
                escaped = f"<sup>{escaped}</sup>"

            if color_int and color_int != 0:
                r = (color_int >> 16) & 0xFF
                g = (color_int >> 8) & 0xFF
                b = color_int & 0xFF
                if not (r < 32 and g < 32 and b < 32):
                    escaped = f'<span style="color:rgb({r},{g},{b})">{escaped}</span>'

            for lk in page_links:
                if lk["rect"].intersects(span_bbox):
                    escaped = f'<a href="{_esc(lk["uri"])}" target="_blank" rel="noopener noreferrer">{escaped}</a>'
                    break

            spans_html.append(escaped)

        line_str = "".join(spans_html).strip()
        if line_str:
            lines_html.append(line_str)

    if not lines_html:
        return ""

    combined = " ".join(lines_html)

    stripped = combined.lstrip()
    bullet_chars = ("•", "●", "○", "▪", "■", "–", "—", "- ", "* ")
    for b in bullet_chars:
        if stripped.startswith(b):
            content = stripped[len(b):].strip()
            return f"<li>{content}</li>"

    ratio = max_font_size / max(median_body_size, 8.0)
    if ratio >= 1.85:
        return f"<h1{align_style}>{combined}</h1>"
    elif ratio >= 1.45:
        return f"<h2{align_style}>{combined}</h2>"
    elif ratio >= 1.22:
        return f"<h3{align_style}>{combined}</h3>"
    elif ratio >= 1.10 and any(k in block.get("lines", [{}])[0].get("spans", [{}])[0].get("font", "").lower() for k in ("bold", "black")):
        return f"<h4{align_style}>{combined}</h4>"

    return f"<p{align_style}>{combined}</p>"


def _render_table_data(table_data: list) -> str:
    """Render a 2D matrix into a clean semantic HTML table."""
    if not table_data:
        return ""
    html = "<table>\n"
    for r_idx, row in enumerate(table_data):
        tag = "th" if r_idx == 0 else "td"
        html += "  <tr>"
        for cell in row:
            val = _esc(str(cell) if cell is not None else "")
            html += f"<{tag}>{val}</{tag}>"
        html += "</tr>\n"
    html += "</table>"
    return html


# ── 2. DOCX Converter (Mammoth + python-docx Fallback) ────────────────────────

def convert_docx(data: bytes, filename: str) -> str:
    import mammoth

    style_map = """
    p[style-name='Title'] => h1.title:fresh
    p[style-name='Subtitle'] => h2.subtitle:fresh
    p[style-name='Heading 1'] => h1:fresh
    p[style-name='Heading 2'] => h2:fresh
    p[style-name='Heading 3'] => h3:fresh
    p[style-name='Heading 4'] => h4:fresh
    p[style-name='Heading 5'] => h5:fresh
    p[style-name='Heading 6'] => h6:fresh
    p[style-name='Quote'] => blockquote > p:fresh
    p[style-name='Intense Quote'] => blockquote > p:fresh
    p[style-name='List Paragraph'] => li:fresh
    r[style-name='Strong'] => strong
    r[style-name='Emphasis'] => em
    r[style-name='Intense Emphasis'] => strong > em
    r[style-name='Subtle Reference'] => cite
    """

    def _docx_img_handler(image):
        with image.open() as stream:
            img_bytes = stream.read()
        ct = image.content_type or "image/png"
        data_uri = _to_base64_data_uri(img_bytes, ct)
        return {"src": data_uri, "style": "max-width:100%;height:auto;border-radius:4px;"}

    try:
        result = mammoth.convert_to_html(
            io.BytesIO(data),
            style_map=style_map,
            convert_image=mammoth.images.img_element(_docx_img_handler),
        )
        inner_html = result.value.replace("<p></p>", "").strip()
        if inner_html:
            return wrap_html(inner_html, filename)
    except Exception:
        pass

    return _convert_docx_via_pythondocx(data, filename)


def _convert_docx_via_pythondocx(data: bytes, filename: str) -> str:
    """Fallback DOCX converter using python-docx."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = []

    rel_images = {}
    for rel_id, rel in doc.part.related_parts.items():
        if "image" in rel.content_type:
            mime = rel.content_type
            b64_uri = _to_base64_data_uri(rel.blob, mime)
            rel_images[rel_id] = b64_uri

    for elem in doc.element.body:
        tag_name = elem.tag.split("}")[-1]

        if tag_name == "p":
            from docx.text.paragraph import Paragraph
            p = Paragraph(elem, doc)
            if not p.text.strip():
                continue

            runs_html = []
            for r in p.runs:
                txt = _esc(r.text)
                if not txt:
                    continue
                if r.bold:
                    txt = f"<strong>{txt}</strong>"
                if r.italic:
                    txt = f"<em>{txt}</em>"
                if r.underline:
                    txt = f"<u>{txt}</u>"
                try:
                    if r.font.color and r.font.color.rgb:
                        rgb = str(r.font.color.rgb)
                        txt = f'<span style="color:#{rgb}">{txt}</span>'
                except Exception:
                    pass
                runs_html.append(txt)

            p_line = "".join(runs_html)
            style_name = p.style.name.lower() if p.style else ""
            if "heading 1" in style_name:
                parts.append(f"<h1>{p_line}</h1>")
            elif "heading 2" in style_name:
                parts.append(f"<h2>{p_line}</h2>")
            elif "heading 3" in style_name:
                parts.append(f"<h3>{p_line}</h3>")
            elif "list" in style_name:
                parts.append(f"<ul><li>{p_line}</li></ul>")
            else:
                parts.append(f"<p>{p_line}</p>")

        elif tag_name == "tbl":
            from docx.table import Table
            t = Table(elem, doc)
            tbl_parts = ['<div class="doc-table-wrap"><table>']
            for ri, row in enumerate(t.rows):
                tag = "th" if ri == 0 else "td"
                tbl_parts.append("  <tr>")
                for cell in row.cells:
                    val = _esc(cell.text.strip())
                    tbl_parts.append(f"    <{tag}>{val}</{tag}>")
                tbl_parts.append("  </tr>")
            tbl_parts.append("</table></div>")
            parts.append("\n".join(tbl_parts))

    for r_id, uri in rel_images.items():
        parts.append(f'<div class="doc-img-wrap"><img src="{uri}" alt="Document Picture"></div>')

    inner = "\n".join(parts) if parts else "<p><em>(Empty document)</em></p>"
    return wrap_html(inner, filename)


# ── 3. PPTX Converter (High-Fidelity Slide Cards) ─────────────────────────────

def convert_pptx(data: bytes, filename: str) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(data))
    slides_html = []

    for s_idx, slide in enumerate(prs.slides, 1):
        parts = []
        title_text = ""

        all_shapes = []
        def collect_shapes(shape_list):
            for shp in shape_list:
                all_shapes.append(shp)
                if shp.shape_type == MSO_SHAPE_TYPE.GROUP and hasattr(shp, "shapes"):
                    collect_shapes(shp.shapes)

        collect_shapes(slide.shapes)
        all_shapes.sort(key=lambda s: (getattr(s, "top", 0) or 0, getattr(s, "left", 0) or 0))

        for shape in all_shapes:
            if shape == slide.shapes.title and shape.has_text_frame:
                title_text = shape.text_frame.text.strip()
                continue

            if shape.has_text_frame:
                txt_html = _render_pptx_text_frame(shape.text_frame)
                if txt_html.strip():
                    parts.append(txt_html)

            if shape.has_table:
                tbl = shape.table
                tbl_rows = []
                for row in tbl.rows:
                    tbl_rows.append([cell.text.strip() for cell in row.cells])
                if tbl_rows:
                    parts.append(f'<div class="doc-table-wrap">{_render_table_data(tbl_rows)}</div>')

            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (hasattr(shape, "image") and shape.image):
                    img_blob = shape.image.blob
                    ct = shape.image.content_type or "image/png"
                    data_uri = _to_base64_data_uri(img_blob, ct)
                    parts.append(f'<div class="doc-img-wrap"><img src="{data_uri}" alt="Slide {s_idx} Graphic"></div>')
            except Exception:
                pass

        notes_html = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    notes_html = (
                        f'<details class="notes-drawer">'
                        f'  <summary>Speaker Notes</summary>'
                        f'  <div class="notes-body">{_esc(notes_text)}</div>'
                        f'</details>'
                    )
        except Exception:
            pass

        slide_title = f'<div class="slide-title">Slide {s_idx}' + (f': {_esc(title_text)}' if title_text else '') + '</div>'
        body = "\n    ".join(parts) if parts else "<p><em>(Empty slide content)</em></p>"
        slides_html.append(
            f'<div class="slide-card">\n'
            f'  {slide_title}\n'
            f'  <div class="slide-body">\n    {body}\n  </div>\n'
            f'  {notes_html}\n'
            f'</div>'
        )

    inner = "\n".join(slides_html) if slides_html else "<p><em>No slides found.</em></p>"
    return wrap_html(inner, filename)


def _render_pptx_text_frame(text_frame) -> str:
    """Render PPTX text frame with per-run formatting, colors, and hyperlinks."""
    paras = []
    for para in text_frame.paragraphs:
        if not para.text.strip():
            continue

        runs_html = []
        for run in para.runs:
            txt = _esc(run.text)
            if not txt:
                continue

            styles = []
            if run.font.bold:
                txt = f"<strong>{txt}</strong>"
            if run.font.italic:
                txt = f"<em>{txt}</em>"
            if run.font.underline:
                txt = f"<u>{txt}</u>"
            if run.font.size and run.font.size.pt:
                styles.append(f"font-size:{run.font.size.pt}pt")

            try:
                if run.font.color and run.font.color.rgb:
                    rgb = str(run.font.color.rgb)
                    styles.append(f"color:#{rgb}")
            except Exception:
                pass

            if styles:
                txt = f'<span style="{";".join(styles)}">{txt}</span>'

            try:
                if run.hyperlink and run.hyperlink.address:
                    txt = f'<a href="{_esc(run.hyperlink.address)}" target="_blank">{txt}</a>'
            except Exception:
                pass

            runs_html.append(txt)

        if runs_html:
            line = "".join(runs_html)
            lvl = para.level or 0
            if lvl > 0:
                paras.append(f'<li style="margin-left:{lvl * 20}px;">{line}</li>')
            else:
                paras.append(f"<p>{line}</p>")

    return "\n".join(paras)


# ── 4. XLSX Converter (Multi-Sheet Tabs & Styles) ─────────────────────────────

def convert_xlsx(data: bytes, filename: str) -> str:
    from openpyxl import load_workbook
    from datetime import datetime, date

    wb = load_workbook(io.BytesIO(data), data_only=True)
    sheet_names = wb.sheetnames
    if not sheet_names:
        return wrap_html("<p><em>Empty workbook.</em></p>", filename)

    sheets_data = []

    for s_idx, s_name in enumerate(sheet_names):
        ws = wb[s_name]
        if ws.max_row is None or ws.max_row == 0:
            continue

        sheet_imgs = []
        try:
            if hasattr(ws, "_images"):
                for im in ws._images:
                    img_bytes = im._data()
                    mime = "image/png"
                    data_uri = _to_base64_data_uri(img_bytes, mime)
                    sheet_imgs.append(f'<div class="doc-img-wrap"><img src="{data_uri}" alt="Sheet Graphic"></div>')
        except Exception:
            pass

        merged = {}
        skip_cells = set()
        for mr in ws.merged_cells.ranges:
            merged[(mr.min_row, mr.min_col)] = (mr.max_row - mr.min_row + 1, mr.max_col - mr.min_col + 1)
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    if (r, c) != (mr.min_row, mr.min_col):
                        skip_cells.add((r, c))

        tbl_parts = ['<div class="doc-table-wrap"><table>']
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column), 1):
            tag = "th" if r_idx == 1 else "td"
            tbl_parts.append("  <tr>")
            for cell in row:
                coord = (cell.row, cell.column)
                if coord in skip_cells:
                    continue

                val_raw = cell.value
                if isinstance(val_raw, (datetime, date)):
                    val = _esc(val_raw.strftime("%Y-%m-%d"))
                elif isinstance(val_raw, float):
                    val = f"{val_raw:,.2f}".rstrip("0").rstrip(".") if "." in f"{val_raw}" else f"{val_raw:,.0f}"
                else:
                    val = _esc(str(val_raw) if val_raw is not None else "")

                styles = []
                try:
                    if cell.font:
                        if cell.font.bold:
                            val = f"<strong>{val}</strong>"
                        if cell.font.color and cell.font.color.rgb and isinstance(cell.font.color.rgb, str):
                            rgb = cell.font.color.rgb[-6:]
                            styles.append(f"color:#{rgb}")
                    if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb and isinstance(cell.fill.start_color.rgb, str):
                        bg = cell.fill.start_color.rgb[-6:]
                        if bg != "000000":
                            styles.append(f"background-color:#{bg}")
                    if cell.alignment and cell.alignment.horizontal:
                        styles.append(f"text-align:{cell.alignment.horizontal}")
                except Exception:
                    pass

                attrs = ""
                if styles:
                    attrs += f' style="{";".join(styles)}"'
                if coord in merged:
                    r_span, c_span = merged[coord]
                    if r_span > 1:
                        attrs += f' rowspan="{r_span}"'
                    if c_span > 1:
                        attrs += f' colspan="{c_span}"'

                tbl_parts.append(f"    <{tag}{attrs}>{val}</{tag}>")
            tbl_parts.append("  </tr>")
        tbl_parts.append("</table></div>")

        full_sheet_html = "\n".join(sheet_imgs) + "\n" + "\n".join(tbl_parts)
        sheets_data.append({"name": s_name, "html": full_sheet_html, "id": f"sheet_{s_idx}"})

    if not sheets_data:
        return wrap_html("<p><em>No data found in workbook.</em></p>", filename)

    if len(sheets_data) == 1:
        return wrap_html(f"<h2>{_esc(sheets_data[0]['name'])}</h2>\n" + sheets_data[0]["html"], filename)

    tab_buttons = []
    tab_panels = []
    for idx, sh in enumerate(sheets_data):
        active_cls = " active" if idx == 0 else ""
        tab_buttons.append(
            f'<button class="tab-btn-sheet{active_cls}" onclick="switchSheetTab(\'{sh["id"]}\')">'
            f'{_esc(sh["name"])}</button>'
        )
        tab_panels.append(
            f'<div id="{sh["id"]}" class="tab-panel-sheet{active_cls}">\n{sh["html"]}\n</div>'
        )

    tab_script = """
    <script>
    function switchSheetTab(sheetId) {
      document.querySelectorAll('.tab-panel-sheet').forEach(p => p.classList.remove('active'));
      document.querySelectorAll('.tab-btn-sheet').forEach(b => b.classList.remove('active'));
      const target = document.getElementById(sheetId);
      if(target) target.classList.add('active');
      event.target.classList.add('active');
    }
    </script>
    """

    inner = (
        f'<div class="tab-container">\n'
        f'  <div class="tab-buttons">\n    {"".join(tab_buttons)}\n  </div>\n'
        f'  {"".join(tab_panels)}\n'
        f'</div>\n'
        f'{tab_script}'
    )
    return wrap_html(inner, filename)


# ── 5. Image Converter (Self-Contained Base64 + Smart OCR) ────────────────────

def convert_image(data: bytes, filename: str) -> str:
    from PIL import Image

    mime = _detect_img_mime(data, filename)
    data_uri = _to_base64_data_uri(data, mime)

    meta_badges = [f'<span class="meta-badge">{_esc(mime)}</span>']
    meta_badges.append(f'<span class="meta-badge">{len(data) // 1024} KB</span>')

    try:
        img = Image.open(io.BytesIO(data))
        meta_badges.append(f'<span class="meta-badge">{img.width} &times; {img.height} px</span>')
        meta_badges.append(f'<span class="meta-badge">{img.mode} mode</span>')
    except Exception:
        pass

    parts = [
        f'<div class="meta-badge-bar">{"".join(meta_badges)}</div>',
        f'<div class="doc-img-wrap">',
        f'  <img src="{data_uri}" alt="{_esc(filename)}" style="max-width:100%;height:auto;border-radius:6px;">',
        f'</div>'
    ]

    ocr_text = _run_ocr_image(data)
    if ocr_text:
        lines = [_esc(l) for l in ocr_text.split("\n") if l.strip()]
        parts.append(
            f'<div class="ocr-box">'
            f'  <div class="ocr-title">Extracted Text (OCR)</div>'
            f'  <div>{"<br>".join(lines)}</div>'
            f'</div>'
        )

    return wrap_html("\n".join(parts), filename)


# ── 6. SVG Converter ──────────────────────────────────────────────────────────

def convert_svg(data: bytes, filename: str) -> str:
    data_uri = _to_base64_data_uri(data, "image/svg+xml")
    body = (
        f'<div class="doc-img-wrap" style="text-align:center;padding:16px;">\n'
        f'  <img src="{data_uri}" alt="{_esc(filename)}" style="max-width:100%;height:auto;">\n'
        f'</div>'
    )
    return wrap_html(body, filename)


# ── 7. Markdown Converter ─────────────────────────────────────────────────────

def convert_markdown(data: bytes, filename: str) -> str:
    import markdown
    text = data.decode("utf-8", errors="replace")
    html = markdown.markdown(
        text,
        extensions=["tables", "fenced_code", "codehilite", "toc", "nl2br", "sane_lists", "attr_list", "def_list"],
    )
    return wrap_html(html, filename)


# ── 8. CSV / TSV Converter ────────────────────────────────────────────────────

def convert_csv(data: bytes, filename: str) -> str:
    text = data.decode("utf-8", errors="replace")
    delim = "\t" if filename.lower().endswith(".tsv") else ","
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
        delim = dialect.delimiter
    except Exception:
        pass

    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = list(reader)
    if not rows:
        return wrap_html("<p><em>Empty file.</em></p>", filename)

    tbl_html = _render_table_data(rows)
    return wrap_html(f'<div class="doc-table-wrap">{tbl_html}</div>', filename)


# ── 9. JSON Converter ─────────────────────────────────────────────────────────

def convert_json(data: bytes, filename: str) -> str:
    text = data.decode("utf-8", errors="replace")
    try:
        parsed = json.loads(text)
        formatted = json.dumps(parsed, indent=2, ensure_ascii=False)
    except Exception:
        formatted = text
    return wrap_html(f'<pre><code>{_esc(formatted)}</code></pre>', filename)


# ── 10. XML Converter ─────────────────────────────────────────────────────────

def convert_xml(data: bytes, filename: str) -> str:
    text = data.decode("utf-8", errors="replace")
    return wrap_html(f'<pre><code>{_esc(text)}</code></pre>', filename)


# ── 11. Plain Text / Code ─────────────────────────────────────────────────────

def convert_text(data: bytes, filename: str) -> str:
    text = data.decode("utf-8", errors="replace")
    paras = []
    for line in text.split("\n"):
        s = line.strip()
        if s:
            paras.append(f"<p>{_esc(s)}</p>")
        else:
            paras.append("<br>")
    return wrap_html("\n".join(paras), filename)


# ── 12. Passthrough ───────────────────────────────────────────────────────────

def convert_passthrough(data: bytes, filename: str) -> str:
    return data.decode("utf-8", errors="replace")


# ─── Converter Registry ───────────────────────────────────────────────────────

CONVERTERS = {
    "pdf": convert_pdf,
    "docx": convert_docx,
    "pptx": convert_pptx,
    "xlsx": convert_xlsx,
    "image": convert_image,
    "svg": convert_svg,
    "markdown": convert_markdown,
    "csv": convert_csv,
    "json": convert_json,
    "xml": convert_xml,
    "text": convert_text,
    "passthrough": convert_passthrough,
}


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# LOCAL TESTING UI (Served when LOCAL_RUN=true)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

LOCAL_TEST_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Document-to-HTML Studio</title>
  <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
  <style>
    *, *::before, *::after { box-sizing: border-box; margin: 0; padding: 0; }
    body {
      font-family: 'Plus Jakarta Sans', system-ui, sans-serif;
      background: #090d16;
      color: #e2e8f0;
      min-height: 100vh;
      display: flex;
      flex-direction: column;
      align-items: center;
      padding: 48px 20px;
    }
    .header { text-align: center; margin-bottom: 36px; max-width: 680px; }
    .header h1 {
      font-size: 2.4rem;
      font-weight: 800;
      background: linear-gradient(135deg, #818cf8 0%, #38bdf8 50%, #34d399 100%);
      -webkit-background-clip: text;
      -webkit-text-fill-color: transparent;
      margin-bottom: 8px;
      letter-spacing: -0.5px;
    }
    .header p { color: #94a3b8; font-size: 1rem; line-height: 1.5; }
    .badge-bar { display: flex; flex-wrap: wrap; gap: 6px; justify-content: center; margin-top: 16px; }
    .badge {
      font-size: 0.72rem; font-weight: 600; padding: 4px 10px;
      border-radius: 999px; background: rgba(129,140,248,0.12);
      color: #a5b4fc; border: 1px solid rgba(129,140,248,0.25);
      text-transform: uppercase; letter-spacing: 0.5px;
    }
    .app-card {
      width: 100%; max-width: 820px;
      background: rgba(17, 24, 39, 0.75);
      backdrop-filter: blur(20px);
      border: 1px solid rgba(255, 255, 255, 0.08);
      border-radius: 20px;
      padding: 36px;
      box-shadow: 0 20px 40px rgba(0, 0, 0, 0.4);
      margin-bottom: 30px;
    }
    .drop-box {
      border: 2px dashed rgba(129,140,248,0.35);
      border-radius: 16px;
      padding: 48px 24px;
      text-align: center;
      cursor: pointer;
      position: relative;
      transition: all 0.25s ease;
      background: rgba(255, 255, 255, 0.02);
    }
    .drop-box:hover, .drop-box.dragover {
      border-color: #38bdf8;
      background: rgba(56,189,248,0.06);
      transform: translateY(-2px);
    }
    .drop-box input[type="file"] { position: absolute; inset: 0; opacity: 0; cursor: pointer; }
    .drop-icon { font-size: 2.8rem; margin-bottom: 12px; display: inline-block; }
    .drop-title { font-size: 1.15rem; font-weight: 700; color: #f8fafc; margin-bottom: 4px; }
    .drop-sub { color: #94a3b8; font-size: 0.88rem; }
    .file-pill {
      margin-top: 16px;
      padding: 12px 18px;
      background: rgba(56,189,248,0.12);
      border: 1px solid rgba(56,189,248,0.25);
      border-radius: 10px;
      font-size: 0.9rem;
      color: #bae6fd;
      display: none;
      align-items: center;
      justify-content: space-between;
    }
    .btn-action {
      display: block; width: 100%; padding: 15px;
      background: linear-gradient(135deg, #6366f1, #3b82f6);
      color: #ffffff; font-weight: 700; font-size: 1rem;
      border: none; border-radius: 12px; cursor: pointer;
      margin-top: 20px; transition: all 0.2s ease;
      box-shadow: 0 4px 16px rgba(99,102,241,0.35);
    }
    .btn-action:hover { transform: translateY(-1px); box-shadow: 0 6px 24px rgba(99,102,241,0.5); }
    .btn-action:disabled { opacity: 0.4; cursor: not-allowed; transform: none; box-shadow: none; }
    .status-msg {
      text-align: center; margin-top: 16px; font-size: 0.9rem; color: #94a3b8; display: none;
    }
    .status-msg.error { color: #f87171; }
    .status-msg.success { color: #34d399; }
    .preview-card {
      width: 100%; max-width: 980px;
      background: rgba(17, 24, 39, 0.75);
      backdrop-filter: blur(20px);
      border: 1px solid rgba(255, 255, 255, 0.08);
      border-radius: 20px;
      padding: 24px;
      display: none;
      box-shadow: 0 20px 40px rgba(0, 0, 0, 0.4);
    }
    .preview-toolbar {
      display: flex; flex-wrap: wrap; align-items: center; justify-content: space-between;
      gap: 12px; margin-bottom: 16px; border-bottom: 1px solid rgba(255,255,255,0.08);
      padding-bottom: 14px;
    }
    .tab-group { display: flex; gap: 4px; background: rgba(0,0,0,0.3); padding: 4px; border-radius: 8px; }
    .tab-btn {
      padding: 6px 14px; border-radius: 6px; border: none; background: transparent;
      color: #94a3b8; font-size: 0.85rem; font-weight: 600; cursor: pointer;
      transition: all 0.15s ease;
    }
    .tab-btn.active { background: #38bdf8; color: #090d16; }
    .action-group { display: flex; gap: 8px; }
    .tool-btn {
      padding: 7px 14px; border-radius: 8px; font-size: 0.82rem; font-weight: 600;
      border: 1px solid rgba(255,255,255,0.12); background: rgba(255,255,255,0.04);
      color: #f1f5f9; cursor: pointer; transition: all 0.2s ease;
    }
    .tool-btn:hover { background: rgba(255,255,255,0.1); }
    .frame-container {
      width: 100%; height: 640px; border: none; border-radius: 12px;
      background: #ffffff;
    }
    .source-container {
      width: 100%; height: 640px; border-radius: 12px;
      background: #0b0f19; color: #a5f3fc; padding: 18px;
      font-family: 'JetBrains Mono', monospace; font-size: 0.85rem;
      overflow: auto; white-space: pre-wrap; display: none;
    }
    .metrics-bar {
      display: flex; gap: 16px; font-size: 0.82rem; color: #94a3b8; margin-top: 14px;
    }
  </style>
</head>
<body>
  <div class="header">
    <h1>Document-to-HTML Studio</h1>
    <p>Convert PDF, Word, PowerPoint, Excel, Images, Markdown & CSV into accurate, self-contained HTML with embedded Base64 images.</p>
    <div class="badge-bar">
      <span class="badge">PDF (PyMuPDF)</span>
      <span class="badge">Word (DOCX)</span>
      <span class="badge">PowerPoint (PPTX)</span>
      <span class="badge">Excel (XLSX)</span>
      <span class="badge">Images (Base64 + OCR)</span>
      <span class="badge">Markdown</span>
      <span class="badge">CSV / TSV</span>
    </div>
  </div>

  <div class="app-card">
    <div class="drop-box" id="dropBox">
      <input type="file" id="fileInput" accept=".pdf,.docx,.doc,.pptx,.ppt,.xlsx,.xls,.png,.jpg,.jpeg,.webp,.bmp,.tiff,.gif,.svg,.md,.csv,.tsv,.json,.xml,.txt,.log,.html">
      <div class="drop-icon">📄</div>
      <div class="drop-title">Drop your document here</div>
      <div class="drop-sub">or click to browse files from your computer</div>
    </div>

    <div class="file-pill" id="filePill">
      <span id="fileNameDisplay">file.pdf</span>
      <span id="fileSizeDisplay" style="opacity:0.7;">0 KB</span>
    </div>

    <button class="btn-action" id="btnConvert" disabled>Convert to HTML</button>
    <div class="status-msg" id="statusMsg"></div>
  </div>

  <div class="preview-card" id="previewCard">
    <div class="preview-toolbar">
      <div class="tab-group">
        <button class="tab-btn active" id="tabPreview">Rendered View</button>
        <button class="tab-btn" id="tabSource">HTML Source</button>
      </div>
      <div class="action-group">
        <button class="tool-btn" id="btnCopy">Copy HTML</button>
        <button class="tool-btn" id="btnDownload">Download .html</button>
      </div>
    </div>

    <iframe class="frame-container" id="previewIframe"></iframe>
    <pre class="source-container" id="sourceView"></pre>

    <div class="metrics-bar" id="metricsBar"></div>
  </div>

  <script>
    const dropBox = document.getElementById('dropBox');
    const fileInput = document.getElementById('fileInput');
    const filePill = document.getElementById('filePill');
    const fileNameDisplay = document.getElementById('fileNameDisplay');
    const fileSizeDisplay = document.getElementById('fileSizeDisplay');
    const btnConvert = document.getElementById('btnConvert');
    const statusMsg = document.getElementById('statusMsg');
    const previewCard = document.getElementById('previewCard');
    const previewIframe = document.getElementById('previewIframe');
    const sourceView = document.getElementById('sourceView');
    const metricsBar = document.getElementById('metricsBar');

    let currentFile = null;
    let convertedHtml = '';

    ['dragenter', 'dragover'].forEach(name => {
      dropBox.addEventListener(name, (e) => { e.preventDefault(); dropBox.classList.add('dragover'); });
    });
    ['dragleave', 'drop'].forEach(name => {
      dropBox.addEventListener(name, (e) => { e.preventDefault(); dropBox.classList.remove('dragover'); });
    });

    dropBox.addEventListener('drop', (e) => {
      if (e.dataTransfer.files && e.dataTransfer.files[0]) {
        handleFileSelect(e.dataTransfer.files[0]);
      }
    });

    fileInput.addEventListener('change', (e) => {
      if (e.target.files && e.target.files[0]) {
        handleFileSelect(e.target.files[0]);
      }
    });

    function handleFileSelect(file) {
      currentFile = file;
      fileNameDisplay.textContent = file.name;
      fileSizeDisplay.textContent = (file.size / 1024).toFixed(1) + ' KB';
      filePill.style.display = 'flex';
      btnConvert.disabled = false;
      statusMsg.style.display = 'none';
    }

    btnConvert.addEventListener('click', async () => {
      if (!currentFile) return;
      btnConvert.disabled = true;
      btnConvert.textContent = 'Converting...';
      statusMsg.className = 'status-msg';
      statusMsg.textContent = 'Processing document accurately...';
      statusMsg.style.display = 'block';

      const startTime = performance.now();
      const formData = new FormData();
      formData.append('file', currentFile);

      try {
        const resp = await fetch('/convert', { method: 'POST', body: formData });
        const data = await resp.json();
        const duration = Math.round(performance.now() - startTime);

        if (!resp.ok || !data.success) {
          throw new Error(data.detail || 'Conversion failed');
        }

        convertedHtml = data.html;
        previewIframe.srcdoc = convertedHtml;
        sourceView.textContent = convertedHtml;

        const imgCount = (convertedHtml.match(/data:image/g) || []).length;
        const tblCount = (convertedHtml.match(/<table/g) || []).length;

        metricsBar.innerHTML = `
          <span>⚡ Time: <strong>${duration} ms</strong></span>
          <span>🖼️ Base64 Images: <strong>${imgCount}</strong></span>
          <span>📊 Tables: <strong>${tblCount}</strong></span>
          <span>📁 Size: <strong>${(data.size_bytes / 1024).toFixed(1)} KB</strong></span>
        `;

        previewCard.style.display = 'block';
        statusMsg.className = 'status-msg success';
        statusMsg.textContent = `Conversion successful in ${duration} ms!`;
        previewCard.scrollIntoView({ behavior: 'smooth' });
      } catch (err) {
        statusMsg.className = 'status-msg error';
        statusMsg.textContent = 'Error: ' + err.message;
      } finally {
        btnConvert.disabled = false;
        btnConvert.textContent = 'Convert to HTML';
      }
    });

    document.getElementById('tabPreview').addEventListener('click', () => {
      document.getElementById('tabPreview').classList.add('active');
      document.getElementById('tabSource').classList.remove('active');
      previewIframe.style.display = 'block';
      sourceView.style.display = 'none';
    });

    document.getElementById('tabSource').addEventListener('click', () => {
      document.getElementById('tabSource').classList.add('active');
      document.getElementById('tabPreview').classList.remove('active');
      sourceView.style.display = 'block';
      previewIframe.style.display = 'none';
    });

    document.getElementById('btnCopy').addEventListener('click', () => {
      if (!convertedHtml) return;
      navigator.clipboard.writeText(convertedHtml).then(() => {
        const btn = document.getElementById('btnCopy');
        btn.textContent = 'Copied!';
        setTimeout(() => btn.textContent = 'Copy HTML', 1500);
      });
    });

    document.getElementById('btnDownload').addEventListener('click', () => {
      if (!convertedHtml) return;
      const blob = new Blob([convertedHtml], { type: 'text/html' });
      const a = document.createElement('a');
      a.href = URL.createObjectURL(blob);
      a.download = (currentFile?.name || 'document').replace(/\\.[^.]+$/, '') + '.html';
      a.click();
    });
  </script>
</body>
</html>"""


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# API ENDPOINTS
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

@app.get("/")
async def root():
    """
    LOCAL_RUN=true  → serves the interactive drag-and-drop web test page
    LOCAL_RUN=false → returns JSON health check (API-only mode)
    """
    if LOCAL_RUN:
        return HTMLResponse(content=LOCAL_TEST_HTML)

    return {
        "service": "Document-to-HTML API",
        "version": "2.1.0",
        "status": "running",
        "docs": "/docs",
        "mode": "api-only",
        "ocr_available": TESSERACT_AVAILABLE,
        "supported_formats": sorted(list(EXTENSION_MAP.keys())),
    }


@app.get("/health")
async def health_check():
    """JSON health-check endpoint regardless of LOCAL_RUN setting."""
    return {
        "status": "healthy",
        "service": "Document-to-HTML API",
        "version": "2.1.0",
        "ocr_available": TESSERACT_AVAILABLE,
        "max_upload_mb": MAX_UPLOAD_BYTES // (1024 * 1024),
    }


@app.post("/convert")
async def convert_document(file: UploadFile = File(...)):
    """
    Upload any document and receive a complete, self-contained HTML document
    with embedded Base64 images and responsive CSS.
    """
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided.")

    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File too large. Maximum size is {MAX_UPLOAD_BYTES // (1024 * 1024)} MB."
        )
    if len(data) == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    ext = Path(file.filename).suffix.lower()
    handler_key = EXTENSION_MAP.get(ext)
    if not handler_key:
        supported = ", ".join(sorted(EXTENSION_MAP.keys()))
        raise HTTPException(
            status_code=415,
            detail=f"Unsupported file extension '{ext}'. Supported: {supported}"
        )

    try:
        html = CONVERTERS[handler_key](data, file.filename)
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Conversion failed: {str(e)}")

    return JSONResponse(content={
        "success": True,
        "filename": file.filename,
        "content_type": file.content_type or "application/octet-stream",
        "size_bytes": len(data),
        "html": html,
    })


@app.post("/convert/raw")
async def convert_document_raw(file: UploadFile = File(...)):
    """
    Upload any document and receive ONLY the inner HTML content
    (without <html>, <head>, or <body> tags) suitable for embedding.
    """
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided.")

    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File exceeds upload limit.")
    if len(data) == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    ext = Path(file.filename).suffix.lower()
    handler_key = EXTENSION_MAP.get(ext)
    if not handler_key:
        raise HTTPException(status_code=415, detail=f"Unsupported format '{ext}'.")

    try:
        html = CONVERTERS[handler_key](data, file.filename)
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Conversion failed: {str(e)}")

    # Strip outer document wrapper to return only inner HTML
    marker_start = '<div class="doc-root">'
    marker_end = "</div>\n</body>"
    i = html.find(marker_start)
    j = html.rfind(marker_end)
    inner = html[i + len(marker_start):j].strip() if i != -1 and j != -1 else html

    return JSONResponse(content={
        "success": True,
        "filename": file.filename,
        "html": inner,
    })


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

if __name__ == "__main__":
    import uvicorn
    print(f"\n  🚀 Document-to-HTML API v2.1")
    print(f"  Mode: {'LOCAL TEST UI at http://localhost:' + str(PORT) if LOCAL_RUN else 'HEADLESS API-ONLY'}")
    print(f"  Port: {PORT}")
    print(f"  Docs: http://localhost:{PORT}/docs")
    print(f"  OCR (Tesseract): {'Available' if TESSERACT_AVAILABLE else 'Not found (optional)'}\n")
    uvicorn.run("app:app", host="0.0.0.0", port=PORT, reload=True)
